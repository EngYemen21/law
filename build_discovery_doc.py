# -*- coding: utf-8 -*-
"""
بناء وثيقة اكتشاف الأعمال (Business Discovery Document) بصيغة DOCX — Arabic RTL.
"""
import sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================================================
# الألوان
# ============================================================
C_PRIMARY = RGBColor(0x0B, 0x1C, 0x2C)
C_BODY    = RGBColor(0x1A, 0x2A, 0x3A)
C_ACCENT  = RGBColor(0x37, 0xDC, 0xF2)
C_ACCENT2 = RGBColor(0x1B, 0x6B, 0x7A)
C_SECOND  = RGBColor(0x50, 0x60, 0x70)
C_SURFACE = RGBColor(0xED, 0xF3, 0xF5)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_GOLD    = RGBColor(0xC9, 0xA8, 0x4C)

ARABIC_FONT = "Arial"
BODY_SIZE = 11
H1_SIZE = 18
H2_SIZE = 14
H3_SIZE = 12

def hexstr(rgb):
    return '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2])

# ============================================================
# XML helpers (RTL + fonts)
# ============================================================
def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)

def set_run_rtl(run):
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)

def set_table_rtl(table):
    tblPr = table._tbl.tblPr
    tblPr.append(OxmlElement('w:bidiVisual'))

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell, color="C8DDE2", size="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)

def no_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        tcBorders.append(el)
    tcPr.append(tcBorders)

# ============================================================
# Content builders
# ============================================================
def add_run(p, text, *, size=BODY_SIZE, bold=False, color=C_BODY, italic=False):
    run = p.add_run(text)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), ARABIC_FONT)
    rFonts.set(qn('w:hAnsi'), ARABIC_FONT)
    rFonts.set(qn('w:cs'), ARABIC_FONT)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size*2)); rPr.append(sz)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(size*2)); rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement('w:b')); rPr.append(OxmlElement('w:bCs'))
    if italic:
        rPr.append(OxmlElement('w:i'))
    if color is not None:
        run.font.color.rgb = color
    set_run_rtl(run)
    return run

def para(doc, text="", *, size=BODY_SIZE, bold=False, color=C_BODY,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0,
         first_line_indent=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before)
    pf.line_spacing = 1.35
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    set_rtl(p)
    if text:
        add_run(p, text, size=size, bold=bold, color=color, italic=italic)
    return p

def bullet(doc, text, *, level=0, bold_prefix=None, color=C_BODY):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.3; pf.space_after = Pt(3)
    pf.right_indent = Cm(0.6 + level*0.6)
    set_rtl(p)
    add_run(p, "◂  ", size=BODY_SIZE, color=C_ACCENT2, bold=True)
    if bold_prefix:
        add_run(p, bold_prefix + ": ", size=BODY_SIZE, bold=True, color=C_PRIMARY)
    add_run(p, text, size=BODY_SIZE, color=color)
    return p

def h1(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    set_rtl(p)
    add_run(p, f"{number}. {text}", size=H1_SIZE, bold=True, color=C_PRIMARY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '6'); bottom.set(qn('w:color'), hexstr(C_ACCENT2))
    pBdr.append(bottom); pPr.append(pBdr)
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_rtl(p)
    add_run(p, text, size=H2_SIZE, bold=True, color=C_ACCENT2)
    return p

def h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_rtl(p)
    add_run(p, text, size=H3_SIZE, bold=True, color=C_PRIMARY)
    return p

def spacer(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    set_rtl(p)
    return p

def page_break(doc):
    p = doc.add_paragraph()
    set_rtl(p)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

def table(doc, headers, rows, *, col_widths_cm=None, header_fill=None):
    if header_fill is None:
        header_fill = hexstr(C_ACCENT2)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(t)
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        set_cell_borders(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        add_run(p, htext, size=BODY_SIZE, bold=True, color=C_WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cells[i], hexstr(C_SURFACE))
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            add_run(p, str(val), size=BODY_SIZE, color=C_BODY)
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

def callout(doc, title, text, *, fill=None, bar=None):
    if fill is None: fill = hexstr(C_SURFACE)
    if bar is None: bar = hexstr(C_ACCENT2)
    t = doc.add_table(rows=1, cols=1)
    set_table_rtl(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell_shading(cell, fill)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','bottom','left'):
        el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4')
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),'FFFFFF'); tcBorders.append(el)
    right = OxmlElement('w:right'); right.set(qn('w:val'),'single'); right.set(qn('w:sz'),'24')
    right.set(qn('w:space'),'0'); right.set(qn('w:color'),bar); tcBorders.append(right)
    tcPr.append(tcBorders)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_rtl(p)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, "◆ " + title, size=BODY_SIZE, bold=True, color=C_ACCENT2)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_rtl(p2)
    p2.paragraph_format.space_before = Pt(2)
    add_run(p2, text, size=BODY_SIZE, color=C_BODY)
    spacer(doc, 6)
    return t

def kv_callout(doc, pairs, *, fill=None, bar=None):
    """صندوق بقائمة نقاط ثابتة."""
    if fill is None: fill = hexstr(C_SURFACE)
    if bar is None: bar = hexstr(C_ACCENT2)
    t = doc.add_table(rows=1, cols=1)
    set_table_rtl(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell_shading(cell, fill)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','bottom','left'):
        el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4')
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),'FFFFFF'); tcBorders.append(el)
    right = OxmlElement('w:right'); right.set(qn('w:val'),'single'); right.set(qn('w:sz'),'24')
    right.set(qn('w:space'),'0'); right.set(qn('w:color'),bar); tcBorders.append(right)
    tcPr.append(tcBorders)
    for k, v in pairs:
        p = cell.add_paragraph() if cell.paragraphs[0].text or len(cell.paragraphs)>1 else cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2)
        add_run(p, k + ": ", size=BODY_SIZE, bold=True, color=C_PRIMARY)
        add_run(p, v, size=BODY_SIZE, color=C_BODY)
    spacer(doc, 6)
    return t

# ============================================================
# Cover page
# ============================================================
def build_cover(doc):
    section = doc.sections[0]
    # لون خلفية الصفحة عبر عنصر background — غير مدعوم بسهولة؛ نستخدم جدول مظلل يملأ الصفحة
    cover_t = doc.add_table(rows=1, cols=1)
    set_table_rtl(cover_t)
    cover_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = cover_t.rows[0].cells[0]
    set_cell_shading(cell, hexstr(C_PRIMARY))
    no_cell_borders(cell)
    # ارتفاع الصف
    tr = cover_t.rows[0]._tr
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), '13800'); trHeight.set(qn('w:hRule'), 'exact')
    trPr.append(trHeight)
    cell.width = Cm(17)

    # محتوى الغلاف
    # مساحة علوية
    p_top = cell.paragraphs[0]
    set_rtl(p_top); p_top.paragraph_format.space_before = Pt(60)

    # شريط لهجة إنجليزي
    p_en = cell.add_paragraph()
    p_en.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_en); p_en.paragraph_format.space_after = Pt(6)
    pPr = p_en._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bb = OxmlElement('w:bottom'); bb.set(qn('w:val'),'single'); bb.set(qn('w:sz'),'6')
    bb.set(qn('w:space'),'8'); bb.set(qn('w:color'), hexstr(C_ACCENT)); pBdr.append(bb)
    pPr.append(pBdr)
    add_run(p_en, "A I - F I R S T   L E G A L   P L A T F O R M", size=10, bold=True, color=C_ACCENT)

    # العنوان الرئيسي
    p_t = cell.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_t); p_t.paragraph_format.space_before = Pt(18); p_t.paragraph_format.space_after = Pt(4)
    add_run(p_t, "وثيقة اكتشاف الأعمال", size=32, bold=True, color=C_WHITE)

    p_t2 = cell.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_t2); p_t2.paragraph_format.space_after = Pt(4)
    add_run(p_t2, "المنصة القانونية المعتمدة على الذكاء الاصطناعي", size=20, bold=True, color=C_ACCENT)

    # عنوان فرعي
    p_sub = cell.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_sub); p_sub.paragraph_format.space_after = Pt(20)
    add_run(p_sub, "نحو نظام تشغيل جديد لمكاتب المحاماة في المملكة العربية السعودية", size=13, color=RGBColor(0xB0,0xB8,0xC0))

    # خط فاصل ذهبي
    p_line = cell.add_paragraph()
    set_rtl(p_line); p_line.paragraph_format.space_after = Pt(14)
    pPr = p_line._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bb = OxmlElement('w:bottom'); bb.set(qn('w:val'),'single'); bb.set(qn('w:sz'),'18')
    bb.set(qn('w:space'),'4'); bb.set(qn('w:color'), hexstr(C_GOLD)); pBdr.append(bb)
    pPr.append(pBdr)

    # شعار افتراضي / اسم المنصة
    p_brand = cell.add_paragraph()
    p_brand.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_brand); p_brand.paragraph_format.space_after = Pt(30)
    add_run(p_brand, "سلاسل بابل القانونية", size=14, bold=True, color=C_GOLD)

    # معلومات وصفية
    for line in [
        "المرحلة: الاكتشاف والتعريف الكامل للأعمال (قبل التصميم التقني)",
        "النطاق: الأعمال — لا تشمل تصميم قاعدة البيانات أو واجهات API أو وكلاء الذكاء الاصطناعي",
        "السوق المستهدف: المملكة العربية السعودية",
        "الإصدار: 1.0",
    ]:
        p_m = cell.add_paragraph()
        p_m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p_m); p_m.paragraph_format.space_after = Pt(4)
        pPr = p_m._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        rb = OxmlElement('w:right'); rb.set(qn('w:val'),'single'); rb.set(qn('w:sz'),'8')
        rb.set(qn('w:space'),'8'); rb.set(qn('w:color'), hexstr(C_ACCENT)); pBdr.append(rb)
        pPr.append(pBdr)
        add_run(p_m, line, size=10, color=RGBColor(0x90,0x98,0x9F))

    # مسافة سفلية + تذييل
    p_sp = cell.add_paragraph()
    set_rtl(p_sp); p_sp.paragraph_format.space_before = Pt(40)

    p_foot = cell.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_foot); p_foot.paragraph_format.space_before = Pt(10)
    pPr = p_foot._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    tb = OxmlElement('w:top'); tb.set(qn('w:val'),'single'); tb.set(qn('w:sz'),'2')
    tb.set(qn('w:space'),'8'); tb.set(qn('w:color'), hexstr(C_ACCENT2)); pBdr.append(tb)
    pPr.append(pBdr)
    add_run(p_foot, "سري — مخصص للاستخدام الداخلي وأصحاب المصلحة المعتمدين", size=9, color=RGBColor(0x68,0x70,0x78))

# ============================================================
# Main
# ============================================================
def main():
    doc = Document()

    # إعداد الصفحة والهوامش للقسم الأول (الغلاف)
    sec0 = doc.sections[0]
    sec0.page_width = Cm(21.0)   # A4
    sec0.page_height = Cm(29.7)
    sec0.top_margin = Cm(0); sec0.bottom_margin = Cm(0)
    sec0.left_margin = Cm(0); sec0.right_margin = Cm(0)
    # RTL للقسم
    sectPr = sec0._sectPr
    bidi = OxmlElement('w:bidi'); sectPr.append(bidi)

    # الغلاف
    build_cover(doc)

    # قسم جديد لمتن الوثيقة
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.top_margin = Cm(2.2); new_sec.bottom_margin = Cm(2.2)
    new_sec.left_margin = Cm(2.0); new_sec.right_margin = Cm(2.0)
    new_sec.page_width = Cm(21.0); new_sec.page_height = Cm(29.7)
    sectPr2 = new_sec._sectPr
    sectPr2.append(OxmlElement('w:bidi'))

    # استيراد محتوى الفصول
    import chapters
    chapters.build_all(doc, globals())

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Business-Discovery-Document.docx")
    doc.save(out)
    print("SAVED:", out)
    print("paragraphs:", len(doc.paragraphs), "| tables:", len(doc.tables))

if __name__ == "__main__":
    main()
