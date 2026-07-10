# -*- coding: utf-8 -*-
"""
بناء تقرير تدقيق معمارية الذكاء الاصطناعي (AI Architecture Audit) — Arabic RTL.
يعيد استخدام المساعدات من build_discovery_doc.py.
"""
import sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import build_discovery_doc as B
from build_discovery_doc import (
    C_PRIMARY, C_BODY, C_ACCENT, C_ACCENT2, C_SECOND, C_SURFACE, C_WHITE, C_GOLD,
    ARABIC_FONT, BODY_SIZE, H1_SIZE, H2_SIZE, H3_SIZE, hexstr,
    set_rtl, set_run_rtl, set_table_rtl, set_cell_shading, set_cell_borders,
    no_cell_borders, add_run, para, bullet, h1, h2, h3, spacer, page_break,
    table as make_table, callout, kv_callout,
)

# ألوان إضافية للنتائج
C_RED    = RGBColor(0xC0, 0x39, 0x2B)
C_AMBER  = RGBColor(0xD4, 0x87, 0x5A)
C_GREEN  = RGBColor(0x2A, 0x6B, 0x4A)

def sev_table(doc, headers, rows, *, col_widths_cm=None, sev_col=None):
    """جدول مع تمييز صفوف حسب الخطورة (خلية في عمود sev_col)."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(t)
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_shading(hdr[i], hexstr(C_ACCENT2))
        set_cell_borders(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        add_run(p, htext, size=BODY_SIZE, bold=True, color=C_WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        sev = row[sev_col] if sev_col is not None else ""
        sev_fill = None
        if sev in ("حرج", "حرجة"): sev_fill = "F8E3E0"
        elif sev in ("عالٍ", "عالية"): sev_fill = "FBF0E6"
        elif sev in ("متوسط", "متوسطة"): sev_fill = hexstr(C_SURFACE)
        elif sev in ("منخفض", "منخفضة"): sev_fill = "E8F0E8"
        for i, val in enumerate(row):
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if sev_fill and i == sev_col:
                set_cell_shading(cells[i], sev_fill)
            elif ridx % 2 == 1:
                set_cell_shading(cells[i], hexstr(C_SURFACE))
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            bold = (i == sev_col and sev in ("حرج","حرجة","عالٍ","عالية"))
            col = C_RED if (i == sev_col and sev in ("حرج","حرجة")) else C_BODY
            add_run(p, str(val), size=BODY_SIZE, color=col, bold=bold)
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

def score_card(doc, label, score, color, comment):
    """بطاقة درجة: خليتان (العدد + الوصف)."""
    t = doc.add_table(rows=1, cols=2)
    set_table_rtl(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = t.rows[0].cells[0]
    set_cell_shading(c1, hexstr(color))
    set_cell_borders(c1, color="FFFFFF", size="4")
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p1)
    p1.paragraph_format.space_before = Pt(8); p1.paragraph_format.space_after = Pt(8)
    add_run(p1, str(score), size=30, bold=True, color=C_WHITE)
    p1b = c1.add_paragraph()
    p1b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p1b)
    add_run(p1b, "/ 100", size=10, color=C_WHITE)
    c1.width = Cm(2.8)
    c2 = t.rows[0].cells[1]
    set_cell_shading(c2, hexstr(C_SURFACE))
    set_cell_borders(c2, color="FFFFFF", size="4")
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p2)
    p2.paragraph_format.space_before = Pt(6); p2.paragraph_format.space_after = Pt(2)
    add_run(p2, label, size=H3_SIZE, bold=True, color=C_PRIMARY)
    p2b = c2.add_paragraph()
    p2b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p2b)
    p2b.paragraph_format.space_after = Pt(6)
    add_run(p2b, comment, size=BODY_SIZE, color=C_BODY)
    c2.width = Cm(11.5)
    spacer(doc, 6)
    return t

# ============================================================
# الغلاف
# ============================================================
def build_cover(doc):
    cover_t = doc.add_table(rows=1, cols=1)
    set_table_rtl(cover_t)
    cover_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = cover_t.rows[0].cells[0]
    set_cell_shading(cell, hexstr(C_PRIMARY))
    no_cell_borders(cell)
    tr = cover_t.rows[0]._tr
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), '13800'); trHeight.set(qn('w:hRule'), 'exact')
    trPr.append(trHeight)
    cell.width = Cm(17)

    p_top = cell.paragraphs[0]
    set_rtl(p_top); p_top.paragraph_format.space_before = Pt(60)

    p_en = cell.add_paragraph()
    p_en.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_en); p_en.paragraph_format.space_after = Pt(6)
    pPr = p_en._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bb = OxmlElement('w:bottom'); bb.set(qn('w:val'),'single'); bb.set(qn('w:sz'),'6')
    bb.set(qn('w:space'),'8'); bb.set(qn('w:color'), hexstr(C_ACCENT)); pBdr.append(bb)
    pPr.append(pBdr)
    add_run(p_en, "A I   A R C H I T E C T U R E   A U D I T", size=10, bold=True, color=C_ACCENT)

    p_t = cell.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_t); p_t.paragraph_format.space_before = Pt(18); p_t.paragraph_format.space_after = Pt(4)
    add_run(p_t, "تقرير تدقيق معمارية الذكاء الاصطناعي", size=30, bold=True, color=C_WHITE)

    p_t2 = cell.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_t2); p_t2.paragraph_format.space_after = Pt(4)
    add_run(p_t2, "تقييم الجاهزية للإنتاج وخارطة الطريق نحو معمارية متعددة الوكلاء", size=17, bold=True, color=C_ACCENT)

    p_sub = cell.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_sub); p_sub.paragraph_format.space_after = Pt(20)
    add_run(p_sub, "منصة سلاسل بابل القانونية — المرحلة الحالية: عميل ذكاء اصطناعي واحد", size=12, color=RGBColor(0xB0,0xB8,0xC0))

    p_line = cell.add_paragraph()
    set_rtl(p_line); p_line.paragraph_format.space_after = Pt(14)
    pPr = p_line._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bb = OxmlElement('w:bottom'); bb.set(qn('w:val'),'single'); bb.set(qn('w:sz'),'18')
    bb.set(qn('w:space'),'4'); bb.set(qn('w:color'), hexstr(C_GOLD)); pBdr.append(bb)
    pPr.append(pBdr)

    p_brand = cell.add_paragraph()
    p_brand.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_brand); p_brand.paragraph_format.space_after = Pt(30)
    add_run(p_brand, "سلاسل بابل القانونية", size=14, bold=True, color=C_GOLD)

    for line in [
        "المهمة: تقييم ما إذا كانت المعمارية الحالية ملائمة لمنصة قانونية معتمدة على الذكاء الاصطناعي للإنتاج",
        "المنهج: تدقيق مبني على الكود الفعلي فقط — لا افتراضات",
        "النطاق: LegalAiService + TicketTriage + طبقة البيانات + التكامل مع المزودين",
        "الإصدار: 1.0",
    ]:
        p_m = cell.add_paragraph()
        p_m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p_m); p_m.paragraph_format.space_after = Pt(4)
        pPr = p_m._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        rb = OxmlElement('w:right'); rb.set(qn('w:val'),'single'); rb.set(qn('w:sz'),'8')
        rb.set(qn('w:space'),'8'); rb.set(qn('w:color'), hexstr(C_ACCENT)); pBdr.append(rb)
        pPr.append(pBdr)
        add_run(p_m, line, size=10, color=RGBColor(0x90,0x98,0x9F))

    p_sp = cell.add_paragraph()
    set_rtl(p_sp); p_sp.paragraph_format.space_before = Pt(40)

    p_foot = cell.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p_foot); p_foot.paragraph_format.space_before = Pt(10)
    pPr = p_foot._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    tb = OxmlElement('w:top'); tb.set(qn('w:val'),'single'); tb.set(qn('w:sz'),'2')
    tb.set(qn('w:space'),'8'); tb.set(qn('w:color'), hexstr(C_ACCENT2)); pBdr.append(tb)
    pPr.append(pBdr)
    add_run(p_foot, "سري — مخصص للفريق التقني وأصحاب المصلحة المعتمدين", size=9, color=RGBColor(0x68,0x70,0x78))

# ============================================================
# Main
# ============================================================
def main():
    doc = Document()
    sec0 = doc.sections[0]
    sec0.page_width = Cm(21.0); sec0.page_height = Cm(29.7)
    sec0.top_margin = Cm(0); sec0.bottom_margin = Cm(0)
    sec0.left_margin = Cm(0); sec0.right_margin = Cm(0)
    sectPr = sec0._sectPr
    sectPr.append(OxmlElement('w:bidi'))

    build_cover(doc)

    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.top_margin = Cm(2.2); new_sec.bottom_margin = Cm(2.2)
    new_sec.left_margin = Cm(2.0); new_sec.right_margin = Cm(2.0)
    new_sec.page_width = Cm(21.0); new_sec.page_height = Cm(29.7)
    sectPr2 = new_sec._sectPr
    sectPr2.append(OxmlElement('w:bidi'))

    import audit_chapters
    audit_chapters.build_all(doc, {
        'para': para, 'bullet': bullet, 'h1': h1, 'h2': h2, 'h3': h3,
        'spacer': spacer, 'page_break': page_break,
        'table': make_table, 'sev_table': sev_table, 'callout': callout,
        'kv_callout': kv_callout, 'score_card': score_card,
        'WD_ALIGN_PARAGRAPH': WD_ALIGN_PARAGRAPH,
        'C_RED': C_RED, 'C_AMBER': C_AMBER, 'C_GREEN': C_GREEN,
        'C_SECOND': C_SECOND, 'C_PRIMARY': C_PRIMARY, 'C_BODY': C_BODY,
        'C_ACCENT': C_ACCENT, 'C_ACCENT2': C_ACCENT2, 'C_SURFACE': C_SURFACE,
        'RGBColor': RGBColor,
    })

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AI-Architecture-Audit.docx")
    doc.save(out)
    print("SAVED:", out)
    print("paragraphs:", len(doc.paragraphs), "| tables:", len(doc.tables))

if __name__ == "__main__":
    main()
